import io
import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

RTL_CHAR_RE = re.compile(r'[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF\uFB50-\uFDFF\uFE70-\uFEFF]')


def is_rtl(text: str) -> bool:
    """Checks if a string contains Urdu, Sindhi, or Arabic characters."""
    if not text:
        return False
    return bool(RTL_CHAR_RE.search(str(text)))


def set_cell_background(cell, fill_hex: str):
    """Sets background color of a Word table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)


def set_paragraph_bidi(p):
    """Sets right-to-left bidirectional property on a paragraph."""
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    bidi.set(qn('w:val'), '1')
    pPr.append(bidi)


def generate_question_paper_docx(
    school_name: str,
    exam_term: str,
    grade: str,
    subject: str,
    time_allowed: str,
    total_marks: str,
    teacher_name: str,
    instructions: str,
    sec_a_title: str,
    sec_a_marks: str,
    sec_a_content: str,
    sec_b_title: str,
    sec_b_marks: str,
    sec_b_content: str,
    sec_c_title: str,
    sec_c_marks: str,
    sec_c_content: str
) -> bytes:
    """
    Generates a professional, fully editable Microsoft Word (.docx) question paper
    with full multilingual layout support for English, Urdu, and Sindhi.
    """
    doc = docx.Document()

    # 1. Page Margins (0.6 inch)
    for sec in doc.sections:
        sec.top_margin = Inches(0.6)
        sec.bottom_margin = Inches(0.6)
        sec.left_margin = Inches(0.6)
        sec.right_margin = Inches(0.6)
        sec.page_width = Inches(8.5)
        sec.page_height = Inches(11.0)

    # 2. Main School Header
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run(school_name.upper())
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(16)
    r_title.font.color.rgb = RGBColor(30, 58, 138)  # #1E3A8A

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(6)
    r_sub = p_sub.add_run(f"EXAMINATION PAPER — {exam_term.upper()}")
    r_sub.bold = True
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(51, 65, 85)  # #334155

    # 3. Metadata Table (2 Columns x 3 Rows)
    meta_table = doc.add_table(rows=3, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    col_widths = [Inches(3.6), Inches(3.6)]
    for row in meta_table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = col_widths[i]
            set_cell_background(cell, "F8FAFC")

    fields = [
        [(f"Grade / Class: {grade}", False), (f"Subject: {subject}", is_rtl(subject))],
        [(f"Time Allowed: {time_allowed}", is_rtl(time_allowed)), (f"Total Marks: {total_marks}", is_rtl(total_marks))],
        [(f"Teacher / Examiner: {teacher_name}", is_rtl(teacher_name)), ("Date: _________________", False)]
    ]

    for r_idx, row_fields in enumerate(fields):
        for c_idx, (field_txt, field_rtl) in enumerate(row_fields):
            cell = meta_table.cell(r_idx, c_idx)
            cell_p = cell.paragraphs[0]
            cell_p.paragraph_format.space_before = Pt(2)
            cell_p.paragraph_format.space_after = Pt(2)

            if field_rtl:
                set_paragraph_bidi(cell_p)
                cell_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r = cell_p.add_run(field_txt)
                r.font.name = "Noto Nastaliq Urdu"
                r.font.size = Pt(10)
            else:
                cell_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                # Split label and value for bold label
                if ":" in field_txt:
                    lbl, val = field_txt.split(":", 1)
                    r_lbl = cell_p.add_run(f"{lbl}:")
                    r_lbl.bold = True
                    r_lbl.font.name = "Calibri"
                    r_lbl.font.size = Pt(9.5)
                    r_val = cell_p.add_run(f" {val.strip()}")
                    r_val.font.name = "Calibri"
                    r_val.font.size = Pt(9.5)
                else:
                    r = cell_p.add_run(field_txt)
                    r.font.name = "Calibri"
                    r.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 4. General Instructions
    if instructions and instructions.strip():
        inst_rtl = is_rtl(instructions)
        p_inst_hdr = doc.add_paragraph()
        p_inst_hdr.paragraph_format.space_before = Pt(4)
        p_inst_hdr.paragraph_format.space_after = Pt(2)

        if inst_rtl:
            set_paragraph_bidi(p_inst_hdr)
            p_inst_hdr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_ih = p_inst_hdr.add_run("ہدایات برائے طلباء:")
            r_ih.bold = True
            r_ih.font.name = "Noto Nastaliq Urdu"
            r_ih.font.size = Pt(11)
        else:
            p_inst_hdr.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r_ih = p_inst_hdr.add_run("General Instructions:")
            r_ih.bold = True
            r_ih.font.name = "Calibri"
            r_ih.font.size = Pt(10)

        for line in instructions.strip().split("\n"):
            line_clean = line.strip()
            if line_clean:
                line_rtl = is_rtl(line_clean)
                p_line = doc.add_paragraph()
                p_line.paragraph_format.space_before = Pt(1)
                p_line.paragraph_format.space_after = Pt(1)

                if line_rtl:
                    set_paragraph_bidi(p_line)
                    p_line.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    r_li = p_line.add_run(f"• {line_clean}")
                    r_li.font.name = "Noto Nastaliq Urdu"
                    r_li.font.size = Pt(10)
                else:
                    p_line.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    r_li = p_line.add_run(f"• {line_clean}")
                    r_li.italic = True
                    r_li.font.name = "Calibri"
                    r_li.font.size = Pt(9.5)
                    r_li.font.color.rgb = RGBColor(71, 85, 105)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Helper function to append section
    def append_docx_section(title, marks, content):
        if not content or not content.strip():
            return

        title_rtl = is_rtl(title)
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.space_before = Pt(10)
        p_sec.paragraph_format.space_after = Pt(4)

        sec_header_text = title.strip()
        if marks and marks.strip():
            sec_header_text += f" ({marks.strip()})"

        if title_rtl:
            set_paragraph_bidi(p_sec)
            p_sec.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r_sec = p_sec.add_run(sec_header_text)
            r_sec.bold = True
            r_sec.font.name = "Noto Nastaliq Urdu"
            r_sec.font.size = Pt(12)
            r_sec.font.color.rgb = RGBColor(30, 58, 138)
        else:
            p_sec.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r_sec = p_sec.add_run(sec_header_text)
            r_sec.bold = True
            r_sec.font.name = "Calibri"
            r_sec.font.size = Pt(12)
            r_sec.font.color.rgb = RGBColor(30, 58, 138)

        # Questions content
        for block in content.strip().split("\n"):
            block_clean = block.strip()
            if block_clean:
                block_rtl = is_rtl(block_clean)
                p_q = doc.add_paragraph()
                p_q.paragraph_format.space_before = Pt(2)
                p_q.paragraph_format.space_after = Pt(2)

                if block_rtl:
                    set_paragraph_bidi(p_q)
                    p_q.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    r_q = p_q.add_run(block_clean)
                    r_q.font.name = "Noto Nastaliq Urdu"
                    r_q.font.size = Pt(11)
                else:
                    p_q.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    r_q = p_q.add_run(block_clean)
                    r_q.font.name = "Calibri"
                    r_q.font.size = Pt(10)
                    r_q.font.color.rgb = RGBColor(15, 23, 42)

    # 5. Append Sections
    append_docx_section(sec_a_title or "Section A", sec_a_marks, sec_a_content)
    append_docx_section(sec_b_title or "Section B", sec_b_marks, sec_b_content)
    append_docx_section(sec_c_title or "Section C", sec_c_marks, sec_c_content)

    buffer = io.BytesIO()
    doc.save(buffer)
    docx_bytes = buffer.getvalue()
    buffer.close()
    return docx_bytes
